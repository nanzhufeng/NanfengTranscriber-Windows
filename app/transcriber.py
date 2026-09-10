from __future__ import annotations

import importlib.util
import gc
import hashlib
import json
import math
import os
import subprocess
import sys
import tempfile
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

from .postprocess import polish_to_simplified_chinese


ProgressCallback = Callable[[dict[str, Any]], None]
CancelCallback = Callable[[], bool]
_NVIDIA_DLL_HANDLES: list[Any] = []
_NVIDIA_DLL_DIRECTORIES_ADDED = False

MODEL_REPOS = {
    "base": "Systran/faster-whisper-base",
    "small": "Systran/faster-whisper-small",
    "medium": "Systran/faster-whisper-medium",
    "large-v3": "Systran/faster-whisper-large-v3",
}


def default_model_cache_dir() -> Path:
    override = os.environ.get("NANFENG_TRANSCRIBER_MODEL_DIR")
    if override:
        return Path(override).expanduser()
    if os.name == "nt" and os.environ.get("LOCALAPPDATA"):
        return Path(os.environ["LOCALAPPDATA"]) / "NanfengTranscriber" / "models"
    cache_home = os.environ.get("XDG_CACHE_HOME")
    if cache_home:
        return Path(cache_home).expanduser() / "nanfeng-transcriber" / "models"
    return Path.home() / ".cache" / "nanfeng-transcriber" / "models"


def _model_cache_marker(cache_dir: Path, model_size: str) -> Path:
    digest = hashlib.sha1(model_size.encode("utf-8", errors="ignore")).hexdigest()[:12]
    return cache_dir / f".ready-{digest}.json"


def _is_frozen_runtime() -> bool:
    return bool(getattr(sys, "frozen", False))


class TranscribeStopped(RuntimeError):
    """用户主动停止当前转写。"""


@dataclass(frozen=True)
class TranscribeOptions:
    output_dir: Path
    model_size: str
    language: str
    compute_mode: str
    export_txt: bool
    export_md: bool
    export_srt: bool
    export_docx: bool
    postprocess_enabled: bool
    ffmpeg_dir: Path | None
    create_video_subfolder: bool = False
    save_beside_video: bool = False


@dataclass(frozen=True)
class TranscribeResult:
    files: list[Path]
    text: str


def resolve_output_directory(source: Path, options: TranscribeOptions) -> Path:
    root = source.parent if options.save_beside_video else options.output_dir
    return root / safe_output_stem(source.stem) if options.create_video_subfolder else root


def default_output_dir() -> Path:
    d_drive = Path("D:/")
    if d_drive.exists():
        return d_drive / "南枫转写"
    return Path.home() / "Downloads" / "南枫转写"


def find_ffmpeg_dir(project_root: Path) -> Path | None:
    candidates = [
        project_root / "tools" / "ffmpeg",
        project_root.parent / "JHlib" / "ffmpeg",
        project_root / "JHlib" / "ffmpeg",
    ]
    for candidate in candidates:
        if (candidate / "ffmpeg.exe").exists() and (candidate / "ffprobe.exe").exists():
            return candidate
    return None


def safe_path_name(text: str, fallback: str = "未命名") -> str:
    import re

    cleaned = re.sub(r'[<>:"/\\|?*\r\n\t]+', " ", text or "").strip()
    cleaned = re.sub(r"\s+", " ", cleaned)
    return (cleaned or fallback)[:120]


def safe_output_stem(text: str, max_chars: int = 72) -> str:
    cleaned = safe_path_name(text)
    if len(cleaned) <= max_chars:
        return cleaned
    digest = hashlib.sha1(cleaned.encode("utf-8", errors="ignore")).hexdigest()[:8]
    return f"{cleaned[: max_chars - 9].rstrip()}_{digest}"


def _raise_if_cancelled(cancel_callback: CancelCallback | None) -> None:
    if cancel_callback and cancel_callback():
        raise TranscribeStopped("用户已停止转写。")


def _ffmpeg_path(options: TranscribeOptions, binary: str) -> str:
    if options.ffmpeg_dir:
        return str(options.ffmpeg_dir / binary)
    return binary


def _subprocess_window_options() -> dict[str, Any]:
    if os.name == "nt" and hasattr(subprocess, "CREATE_NO_WINDOW"):
        return {"creationflags": subprocess.CREATE_NO_WINDOW}
    return {}


def probe_duration_seconds(source: Path, options: TranscribeOptions) -> float:
    command = [
        _ffmpeg_path(options, "ffprobe.exe"),
        "-v",
        "error",
        "-show_entries",
        "format=duration",
        "-of",
        "default=noprint_wrappers=1:nokey=1",
        str(source),
    ]
    try:
        completed = subprocess.run(command, capture_output=True, text=True, check=True, **_subprocess_window_options())
        return max(0.0, float(completed.stdout.strip()))
    except Exception:
        return 0.0


def _format_timestamp(seconds: float) -> str:
    seconds = max(0.0, seconds)
    total_ms = int(round(seconds * 1000))
    total_seconds, ms = divmod(total_ms, 1000)
    minutes, sec = divmod(total_seconds, 60)
    hours, minutes = divmod(minutes, 60)
    return f"{hours:02d}:{minutes:02d}:{sec:02d}.{ms:03d}"


def _format_srt_timestamp(seconds: float) -> str:
    return _format_timestamp(seconds).replace(".", ",")


def _format_eta(seconds: float) -> str:
    if not math.isfinite(seconds) or seconds <= 0:
        return "-"
    minutes, sec = divmod(int(seconds), 60)
    hours, minutes = divmod(minutes, 60)
    if hours:
        return f"{hours}时{minutes:02d}分"
    if minutes:
        return f"{minutes}分{sec:02d}秒"
    return f"{sec}秒"


def _clean_segment_text(text: str) -> str:
    return " ".join((text or "").strip().split())


def _should_merge_segments(current: dict[str, Any], next_segment: dict[str, Any], max_gap: float, max_chars: int) -> bool:
    current_text = current["text"].strip()
    next_text = next_segment["text"].strip()
    if not current_text or not next_text:
        return True
    if current_text.endswith(("。", "！", "？", "!", "?", "；", ";")):
        return False
    gap = float(next_segment["start"]) - float(current["end"])
    if gap > max_gap:
        return False
    return len(current_text) + len(next_text) <= max_chars


def _join_segment_text(left: str, right: str) -> str:
    if not left:
        return right
    if not right:
        return left
    if left.endswith(("，", "、", "。", "！", "？", "；", ",", ".", "!", "?", ";")):
        return left + right
    return left + "，" + right


def _merge_readable_segments(segments: list[dict[str, Any]], max_gap: float = 1.0, max_chars: int = 90) -> list[dict[str, Any]]:
    # TXT / Markdown 更偏阅读，合并过碎的 Whisper segment；SRT 仍保留原始字幕节奏。
    merged: list[dict[str, Any]] = []
    for segment in segments:
        text = _clean_segment_text(segment.get("text", ""))
        if not text:
            continue
        item = {"start": segment["start"], "end": segment["end"], "text": text}
        if merged and _should_merge_segments(merged[-1], item, max_gap, max_chars):
            merged[-1]["end"] = item["end"]
            merged[-1]["text"] = _join_segment_text(merged[-1]["text"], item["text"])
        else:
            merged.append(item)
    return merged


def _extract_audio(source: Path, target: Path, options: TranscribeOptions, cancel_callback: CancelCallback | None) -> None:
    command = [
        _ffmpeg_path(options, "ffmpeg.exe"),
        "-y",
        "-v",
        "error",
        "-i",
        str(source),
        "-vn",
        "-ac",
        "1",
        "-ar",
        "16000",
        "-f",
        "wav",
        str(target),
    ]
    process = subprocess.Popen(command, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE, **_subprocess_window_options())
    try:
        while True:
            return_code = process.poll()
            if return_code is not None:
                if return_code != 0:
                    _stdout, stderr = process.communicate()
                    detail = (stderr or b"").decode("utf-8", errors="replace").strip().lower()
                    if "does not contain any stream" in detail or "output file #0 does not contain any stream" in detail:
                        raise RuntimeError("未检测到可转写的音轨。该视频可能没有声音或只有画面。")
                    if "no such file or directory" in detail:
                        raise RuntimeError("音频提取失败：源文件不存在或已被移动。")
                    if "permission denied" in detail:
                        raise RuntimeError("音频提取失败：无法读取源文件，请检查文件权限。")
                    if "invalid data found" in detail:
                        raise RuntimeError("音频提取失败：媒体文件损坏或编码不受支持。")
                    raise RuntimeError("音频提取失败：FFmpeg 无法读取该媒体文件。")
                return
            _raise_if_cancelled(cancel_callback)
            time.sleep(0.2)
    except TranscribeStopped:
        process.terminate()
        try:
            process.wait(timeout=3)
        except subprocess.TimeoutExpired:
            process.kill()
            process.wait(timeout=3)
        raise


def _write_txt(target: Path, segments: list[dict[str, Any]]) -> None:
    readable_segments = _merge_readable_segments(segments)
    lines = [segment["text"].strip() for segment in readable_segments if segment["text"].strip()]
    target.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_md(target: Path, source: Path, segments: list[dict[str, Any]]) -> None:
    readable_segments = _merge_readable_segments(segments)
    lines = [
        f"# {source.stem}",
        "",
        "| 时间 | 内容 |",
        "| --- | --- |",
    ]
    for segment in readable_segments:
        text = segment["text"].strip().replace("|", "\\|")
        if not text:
            continue
        lines.append(f"| {_format_timestamp(segment['start'])} | {text} |")
    target.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_md_text(target: Path, source: Path, text: str) -> None:
    lines = [f"# {source.stem}", ""]
    for paragraph in text.splitlines():
        paragraph = paragraph.strip()
        if paragraph:
            lines.append(paragraph)
            lines.append("")
    target.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def _subtitle_cues(segments: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Build short sequential cues using word alignment when available."""
    cues: list[dict[str, Any]] = []
    for segment in segments:
        units: list[dict[str, Any]] = []
        for word in segment.get("words") or [segment]:
            text = " ".join(word["text"].split())
            if not text:
                continue
            start = max(0.0, float(word["start"]))
            end = float(word["end"])
            if end <= start:
                continue
            # Missing word alignment: distribute only within the original interval.
            count = max(1, math.ceil(len(text) / 24), math.ceil((end - start) / 5))
            count = min(count, len(text))
            for i in range(count):
                lo, hi = i * len(text) // count, (i + 1) * len(text) // count
                units.append({"start": start + (end - start) * lo / len(text),
                              "end": start + (end - start) * hi / len(text),
                              "text": text[lo:hi], "space": str(word["text"]).startswith(" ") and i == 0})
        current = None
        for unit in units:
            separator = " " if unit["space"] else ""
            if current and (len(current["text"] + separator + unit["text"]) > 24
                            or unit["end"] - current["start"] > 5
                            or unit["start"] - current["end"] > 0.6):
                cues.append(current)
                current = None
            if current is None:
                current = {"start": unit["start"], "end": unit["end"], "text": unit["text"]}
            else:
                current["text"] += separator + unit["text"]
                current["end"] = unit["end"]
            if current["text"].endswith(tuple("。！？.!?；;")):
                cues.append(current)
                current = None
        if current:
            cues.append(current)
    cues.sort(key=lambda cue: cue["start"])
    for index in range(len(cues) - 1):
        cues[index]["end"] = min(cues[index]["end"], cues[index + 1]["start"])
    return [cue for cue in cues if round(cue["end"] * 1000) > round(cue["start"] * 1000)]


def _write_srt(target: Path, segments: list[dict[str, Any]]) -> None:
    lines: list[str] = []
    index = 1
    for segment in _subtitle_cues(segments):
        text = segment["text"].strip()
        if not text:
            continue
        lines.extend(
            [
                str(index),
                f"{_format_srt_timestamp(segment['start'])} --> {_format_srt_timestamp(segment['end'])}",
                text,
                "",
            ]
        )
        index += 1
    target.write_text("\n".join(lines), encoding="utf-8")


def _write_docx(target: Path, source: Path, text: str) -> None:
    if importlib.util.find_spec("docx") is None:
        raise RuntimeError("缺少 python-docx 依赖。请重新运行依赖安装后再导出 DOCX。")
    from docx import Document

    document = Document()
    document.add_heading(source.stem, level=1)
    for paragraph in text.splitlines():
        paragraph = paragraph.strip()
        if paragraph:
            document.add_paragraph(paragraph)
    document.save(str(target))


def _dll_exists_in_path(dll_names: set[str]) -> bool:
    if not _is_frozen_runtime():
        for package_name in ("nvidia.cudnn", "nvidia.cublas", "nvidia.cuda_nvrtc"):
            spec = importlib.util.find_spec(package_name)
            if spec and spec.submodule_search_locations:
                bin_dir = Path(list(spec.submodule_search_locations)[0]) / "bin"
                for dll_name in dll_names:
                    if (bin_dir / dll_name).exists():
                        return True
    for folder in os.environ.get("PATH", "").split(os.pathsep):
        if not folder:
            continue
        base = Path(folder)
        for dll_name in dll_names:
            if (base / dll_name).exists():
                return True
    return False


def _has_windows_cuda_runtime() -> bool:
    if os.name != "nt":
        return True
    return _dll_exists_in_path({"cublas64_12.dll"}) and _dll_exists_in_path({"cudnn64_9.dll"})


def _add_nvidia_dll_directories() -> None:
    global _NVIDIA_DLL_DIRECTORIES_ADDED
    if os.name != "nt" or not hasattr(os, "add_dll_directory"):
        return
    if _NVIDIA_DLL_DIRECTORIES_ADDED:
        return
    if _is_frozen_runtime():
        _NVIDIA_DLL_DIRECTORIES_ADDED = True
        return
    for package_name in ("nvidia.cudnn", "nvidia.cublas", "nvidia.cuda_nvrtc"):
        spec = importlib.util.find_spec(package_name)
        if not spec or not spec.submodule_search_locations:
            continue
        bin_dir = Path(list(spec.submodule_search_locations)[0]) / "bin"
        if not bin_dir.exists():
            continue
        _NVIDIA_DLL_HANDLES.append(os.add_dll_directory(str(bin_dir)))
    _NVIDIA_DLL_DIRECTORIES_ADDED = True


def _is_gpu_runtime_error(error_text: str) -> bool:
    message = error_text.lower()
    return any(keyword in message for keyword in ("cuda", "cublas", "cudnn", "cudart", "gpu"))


def _is_gpu_batch_resource_error(error_text: str) -> bool:
    message = error_text.lower()
    return any(
        keyword in message
        for keyword in (
            "out of memory",
            "cuda_error_out_of_memory",
            "failed to allocate",
            "cublas_status_alloc_failed",
            "allocation failed",
        )
    )


@dataclass
class LoadedRuntime:
    model: Any
    pipeline: Any | None
    device: str
    compute_type: str


class TranscriptionSession:
    """复用单次队列中已经加载的转写模型与 GPU 推理管线。"""

    def __init__(self, model_class: Any | None = None, batched_pipeline_class: Any | None = None) -> None:
        self._model_class = model_class
        self._batched_pipeline_class = batched_pipeline_class
        self._runtimes: dict[tuple[str, str, str], LoadedRuntime] = {}
        self._preferred_batch_sizes: dict[tuple[str, str, str], int] = {}
        self._force_cpu = False

    def _backend_classes(self) -> tuple[Any, Any]:
        if self._model_class is None or self._batched_pipeline_class is None:
            from faster_whisper import BatchedInferencePipeline, WhisperModel

            self._model_class = self._model_class or WhisperModel
            self._batched_pipeline_class = self._batched_pipeline_class or BatchedInferencePipeline
        return self._model_class, self._batched_pipeline_class

    def get_runtime(self, options: TranscribeOptions, progress_callback: ProgressCallback) -> LoadedRuntime:
        if options.compute_mode == "CPU 稳定" or self._force_cpu:
            return self._get_or_create_runtime(options, progress_callback, "cpu", "int8")

        if not _has_windows_cuda_runtime():
            self._force_cpu = True
            progress_callback(
                {
                    "status": "loading_model",
                    "progress": 3,
                    "eta": "已切换 CPU",
                    "notice": "未检测到完整 CUDA/cuBLAS/cuDNN 运行库，已跳过 GPU 并切换 CPU。",
                    "force_cpu": True,
                }
            )
            return self._get_or_create_runtime(options, progress_callback, "cpu", "int8")

        try:
            return self._get_or_create_runtime(options, progress_callback, "cuda", "float16")
        except Exception as exc:
            if not _is_gpu_runtime_error(str(exc)):
                raise
            self._force_cpu = True
            progress_callback(
                {
                    "status": "loading_model",
                    "progress": 3,
                    "eta": "GPU 不可用，已切换 CPU",
                    "notice": "GPU 运行库加载失败，已自动改用 CPU 继续转写。",
                    "force_cpu": True,
                }
            )
            return self._get_or_create_runtime(options, progress_callback, "cpu", "int8")

    def _get_or_create_runtime(
        self,
        options: TranscribeOptions,
        progress_callback: ProgressCallback,
        device: str,
        compute_type: str,
    ) -> LoadedRuntime:
        cache_key = (options.model_size, device, compute_type)
        cached = self._runtimes.get(cache_key)
        if cached is not None:
            return cached

        model_class, pipeline_class = self._backend_classes()
        target_name = "GPU" if device == "cuda" else "CPU"
        cache_dir = default_model_cache_dir()
        cache_dir.mkdir(parents=True, exist_ok=True)
        marker = _model_cache_marker(cache_dir, options.model_size)
        cache_ready = marker.is_file()
        progress_callback(
            {
                "status": "loading_model" if cache_ready else "downloading_model",
                "progress": 3,
                "eta": f"正在加载 {target_name} 模型" if cache_ready else f"首次下载 {options.model_size} 模型",
                "notice": (
                    f"正在从本地缓存加载 {options.model_size} 模型。"
                    if cache_ready
                    else f"首次使用需要下载 {options.model_size} 模型，完成后将长期保存在本机。"
                ),
            }
        )
        model_kwargs = {
            "device": device,
            "compute_type": compute_type,
            "download_root": str(cache_dir),
            "local_files_only": cache_ready,
        }
        try:
            model = model_class(options.model_size, **model_kwargs)
        except Exception:
            if not cache_ready:
                raise
            marker.unlink(missing_ok=True)
            progress_callback(
                {
                    "status": "downloading_model",
                    "progress": 3,
                    "eta": f"正在修复 {options.model_size} 模型缓存",
                    "notice": "本地模型缓存不完整，正在联网修复；修复完成后仍会长期复用。",
                }
            )
            model_kwargs["local_files_only"] = False
            model = model_class(options.model_size, **model_kwargs)
        marker.write_text(
            json.dumps(
                {
                    "model": options.model_size,
                    "repo": MODEL_REPOS.get(options.model_size, options.model_size),
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        pipeline = pipeline_class(model=model) if device == "cuda" else None
        runtime = LoadedRuntime(model=model, pipeline=pipeline, device=device, compute_type=compute_type)
        self._runtimes[cache_key] = runtime
        return runtime

    def transcribe_audio(
        self,
        audio_path: Path,
        duration: float,
        options: TranscribeOptions,
        progress_callback: ProgressCallback,
        cancel_callback: CancelCallback | None = None,
    ) -> list[dict[str, Any]]:
        _raise_if_cancelled(cancel_callback)
        runtime = self.get_runtime(options, progress_callback)
        if runtime.device == "cuda" and runtime.pipeline is not None:
            transcribe_kwargs = self._build_transcribe_kwargs(options, include_initial_prompt=False)
            return self._transcribe_with_gpu_batches(
                audio_path,
                duration,
                options,
                runtime,
                transcribe_kwargs,
                progress_callback,
                cancel_callback,
            )
        transcribe_kwargs = self._build_transcribe_kwargs(options)
        return self._run_model_transcription(
            runtime.model,
            audio_path,
            duration,
            transcribe_kwargs,
            progress_callback,
            cancel_callback,
        )

    def _build_transcribe_kwargs(
        self,
        options: TranscribeOptions,
        include_initial_prompt: bool = True,
    ) -> dict[str, Any]:
        language = _resolve_language(options.language)
        transcribe_kwargs: dict[str, Any] = {
            "language": language,
            "vad_filter": True,
            "word_timestamps": options.export_srt,
            "beam_size": 3,
            "temperature": 0.0,
            "condition_on_previous_text": False,
        }
        if language == "zh" and include_initial_prompt:
            transcribe_kwargs["initial_prompt"] = (
                "以下是音视频中的口播内容。请使用现代简体中文输出；"
                "如果听到英文、繁体字表达或文言表达，请尽量转成自然、通顺的简体中文意思。"
            )
        return transcribe_kwargs

    def _transcribe_with_gpu_batches(
        self,
        audio_path: Path,
        duration: float,
        options: TranscribeOptions,
        runtime: LoadedRuntime,
        transcribe_kwargs: dict[str, Any],
        progress_callback: ProgressCallback,
        cancel_callback: CancelCallback | None,
    ) -> list[dict[str, Any]]:
        cache_key = (options.model_size, runtime.device, runtime.compute_type)
        preferred_size = self._preferred_batch_sizes.get(cache_key, 8)
        batch_sizes = [size for size in (8, 4, 2) if size <= preferred_size]

        for batch_size in batch_sizes:
            try:
                segment_iter, _info = runtime.pipeline.transcribe(
                    str(audio_path),
                    batch_size=batch_size,
                    **transcribe_kwargs,
                )
                segments = self._collect_segments(
                    segment_iter,
                    duration,
                    progress_callback,
                    cancel_callback,
                )
            except TranscribeStopped:
                raise
            except Exception as exc:
                error_text = str(exc)
                if _is_gpu_batch_resource_error(error_text):
                    progress_callback(
                        {
                            "status": "loading_model",
                            "progress": 3,
                            "eta": "正在降低 GPU 批量大小",
                            "notice": f"GPU 批量大小 {batch_size} 不可用，正在尝试更低档位。",
                        }
                    )
                    continue
                if _is_gpu_runtime_error(error_text):
                    return self._transcribe_with_cpu_fallback(
                        audio_path,
                        duration,
                        options,
                        progress_callback,
                        cancel_callback,
                    )
                raise
            else:
                self._preferred_batch_sizes[cache_key] = batch_size
                return segments

        try:
            normal_gpu_kwargs = self._build_transcribe_kwargs(options)
            return self._run_model_transcription(
                runtime.model,
                audio_path,
                duration,
                normal_gpu_kwargs,
                progress_callback,
                cancel_callback,
            )
        except TranscribeStopped:
            raise
        except Exception as exc:
            if not _is_gpu_runtime_error(str(exc)):
                raise
            return self._transcribe_with_cpu_fallback(
                audio_path,
                duration,
                options,
                progress_callback,
                cancel_callback,
            )

    def _transcribe_with_cpu_fallback(
        self,
        audio_path: Path,
        duration: float,
        options: TranscribeOptions,
        progress_callback: ProgressCallback,
        cancel_callback: CancelCallback | None,
    ) -> list[dict[str, Any]]:
        self._force_cpu = True
        progress_callback(
            {
                "status": "loading_model",
                "progress": 3,
                "eta": "GPU 不可用，CPU 重试",
                "notice": "GPU 推理失败，已改用 CPU 继续转写；后续任务也会使用 CPU。",
                "force_cpu": True,
            }
        )
        runtime = self.get_runtime(options, progress_callback)
        transcribe_kwargs = self._build_transcribe_kwargs(options)
        return self._run_model_transcription(
            runtime.model,
            audio_path,
            duration,
            transcribe_kwargs,
            progress_callback,
            cancel_callback,
        )

    def _run_model_transcription(
        self,
        model: Any,
        audio_path: Path,
        duration: float,
        transcribe_kwargs: dict[str, Any],
        progress_callback: ProgressCallback,
        cancel_callback: CancelCallback | None,
    ) -> list[dict[str, Any]]:
        segment_iter, _info = model.transcribe(str(audio_path), **transcribe_kwargs)
        return self._collect_segments(segment_iter, duration, progress_callback, cancel_callback)

    def _collect_segments(
        self,
        segment_iter: Any,
        duration: float,
        progress_callback: ProgressCallback,
        cancel_callback: CancelCallback | None,
    ) -> list[dict[str, Any]]:
        started_at = time.monotonic()
        segments: list[dict[str, Any]] = []
        for segment in segment_iter:
            _raise_if_cancelled(cancel_callback)
            entry = {"start": segment.start, "end": segment.end, "text": segment.text}
            words = getattr(segment, "words", None)
            if words:
                entry["words"] = [{"start": w.start, "end": w.end, "text": w.word} for w in words]
            segments.append(entry)
            ratio = min(0.98, (segment.end / duration) if duration else 0.5)
            elapsed = max(time.monotonic() - started_at, 0.1)
            eta = elapsed * (1 - ratio) / ratio if ratio > 0 else 0
            progress_callback(
                {
                    "status": "transcribing",
                    "progress": int(ratio * 100),
                    "eta": _format_eta(eta),
                }
            )
        return segments

    def close(self) -> None:
        self._runtimes.clear()
        self._preferred_batch_sizes.clear()
        gc.collect()


def _resolve_language(language: str) -> str | None:
    language_map = {
        "中文": "zh",
        "英文": "en",
        "日语": "ja",
        "韩语": "ko",
        "zh": "zh",
        "en": "en",
        "ja": "ja",
        "ko": "ko",
    }
    if language == "自动识别":
        return None
    return language_map.get(language, language)


def transcribe_file(
    source: Path,
    options: TranscribeOptions,
    progress_callback: ProgressCallback,
    cancel_callback: CancelCallback | None = None,
    session: TranscriptionSession | None = None,
) -> TranscribeResult:
    if importlib.util.find_spec("faster_whisper") is None:
        raise RuntimeError(
            "缺少 faster-whisper 转写依赖。\n\n"
            "处理方式：在项目目录运行：\n"
            "pip install -r requirements.txt\n\n"
            "安装后重新打开 bat 再开始转写。"
        )

    _add_nvidia_dll_directories()

    _raise_if_cancelled(cancel_callback)
    resolve_output_directory(source, options).mkdir(parents=True, exist_ok=True)
    duration = probe_duration_seconds(source, options)
    owns_session = session is None
    active_session = session or TranscriptionSession()

    try:
        progress_callback({"status": "extracting", "progress": 0, "eta": "-"})
        with tempfile.TemporaryDirectory(prefix="nanzhufeng_transcribe_") as temp_dir:
            audio_path = Path(temp_dir) / "audio.wav"
            _extract_audio(source, audio_path, options, cancel_callback)
            _raise_if_cancelled(cancel_callback)
            segments = active_session.transcribe_audio(
                audio_path,
                duration,
                options,
                progress_callback,
                cancel_callback,
            )
    finally:
        if owns_session:
            active_session.close()

    if not segments:
        raise RuntimeError("没有识别到可导出的文字。")

    safe_stem = safe_output_stem(source.stem)
    output_base = resolve_output_directory(source, options)
    output_base.mkdir(parents=True, exist_ok=True)
    files: list[Path] = []
    readable_segments = _merge_readable_segments(segments)
    plain_text = "\n".join(segment["text"].strip() for segment in readable_segments if segment["text"].strip())
    final_text = plain_text
    if options.postprocess_enabled:
        final_text = polish_to_simplified_chinese(plain_text, source.stem, progress_callback, cancel_callback)
    if options.export_txt:
        txt_path = output_base / f"{safe_stem}.txt"
        if options.postprocess_enabled:
            txt_path.write_text(final_text.rstrip() + "\n", encoding="utf-8")
        else:
            _write_txt(txt_path, segments)
        files.append(txt_path)
    if options.export_md:
        md_path = output_base / f"{safe_stem}.md"
        if options.postprocess_enabled:
            _write_md_text(md_path, source, final_text)
        else:
            _write_md(md_path, source, segments)
        files.append(md_path)
    if options.export_srt:
        srt_path = output_base / f"{safe_stem}.srt"
        _write_srt(srt_path, segments)
        files.append(srt_path)
    if options.export_docx:
        docx_path = output_base / f"{safe_stem}.docx"
        _write_docx(docx_path, source, final_text)
        files.append(docx_path)

    progress_callback({"status": "finished", "progress": 100, "eta": "0秒"})
    return TranscribeResult(files=files, text=final_text)
